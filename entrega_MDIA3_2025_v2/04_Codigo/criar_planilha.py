#!/usr/bin/env python3
from pathlib import Path
from copy import copy
import csv, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.pagebreak import Break
from openpyxl.formatting.rule import CellIsRule

SCRIPT_DIR=Path(__file__).resolve().parent
PROJECT=SCRIPT_DIR.parent if SCRIPT_DIR.name=='04_Codigo' else SCRIPT_DIR
DATA_DIR=PROJECT/'02_Base_e_amostras' if (PROJECT/'02_Base_e_amostras').exists() else PROJECT
SHEET_DIR=PROJECT/'03_Planilhas' if (PROJECT/'03_Planilhas').exists() else PROJECT
MATRIX=PROJECT/'05_Fontes_e_normas'/'matriz_cobertura_normas.csv'
if not MATRIX.exists(): MATRIX=PROJECT/'matriz_cobertura_normas.csv'
OUT=SHEET_DIR/'planilha_auditoria_MDI.xlsx'
DARK='135B44'; WHITE='FFFFFF'; LIGHT='CFE9E0'; GRAY='E7E5E4'; BLUE='0000FF'; BLACK='000000'; GREEN='008000'
source='Fonte: M. Dias Branco, Demonstrações Financeiras Anuais Completas 2025/DFP, exercício findo em 31/12/2025, consolidado, valores em BRL mil; saldo conforme balanço/nota identificada.'
wb=Workbook(); ws=wb.active; ws.title='Leia-me'

def setup(s,title,subtitle,unit):
    s.sheet_view.showGridLines=False
    s.column_dimensions['A'].width=20; s.column_dimensions['B'].width=20
    s['C3']=title; s['C3'].fill=PatternFill('solid',fgColor=DARK); s['C3'].font=Font(name='Arial',size=16,bold=True,color=WHITE)
    s['C5']=subtitle; s['C5'].font=Font(name='Arial',size=11,bold=True)
    s['C6']=unit; s['C6'].font=Font(name='Arial',size=10,italic=True)
    s.row_dimensions[3].height=25
    s.freeze_panes='C8'
    s.page_setup.orientation='landscape'; s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0; s.sheet_properties.pageSetUpPr.fitToPage=True
    s.print_area='B2:K60'; s.oddFooter.center.text='&A | Página &P de &N'

def table(s, headers, rows, startrow=8, startcol=3, name='T'):
    for j,h in enumerate(headers,startcol):
        c=s.cell(startrow,j,h); c.font=Font(bold=True,color=WHITE); c.fill=PatternFill('solid',fgColor=DARK); c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,row in enumerate(rows,startrow+1):
        for j,val in enumerate(row,startcol):
            c=s.cell(i,j,val); c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(val,str) and val.startswith('='): c.font=Font(color=BLACK)
            elif isinstance(val,(int,float)): c.font=Font(color=BLUE); c.comment=Comment(source,'Fonte')
            else: c.font=Font(color=BLACK)
    s.auto_filter.ref=f'{get_column_letter(startcol)}{startrow}:{get_column_letter(startcol+len(headers)-1)}{startrow+len(rows)}'
    for j in range(startcol,startcol+len(headers)):
        letter=get_column_letter(j)
        maxlen=max([len(str(s.cell(i,j).value or '')) for i in range(startrow,startrow+len(rows)+1)] or [12])
        s.column_dimensions[letter].width=min(max(14,maxlen+2),46)
    s.row_dimensions[startrow].height=32
    for i in range(startrow+1,startrow+len(rows)+1): s.row_dimensions[i].height=35

setup(ws,'Guia de uso','Reconstrução acadêmica da auditoria — matrícula 202062044','Dados públicos e simulação sintética | sem opinião independente')
intro=[
('Identificação','Matrícula 202062044.'),
('Período','Exercício encerrado em 31/12/2025; relatório publicado em 26/02/2026.'),
('Unidade','BRL mil, salvo indicação em contrário.'),
('Escopo','Base e exceções são sintéticas; apenas saldos agregados da Nota 11 foram calibrados.'),
('Planilha','Contém saldos, materialidade, riscos, análises, programas, amostras, exceções e cobertura de normas.'),
('Limite','Os testes computacionais usam dados sintéticos; não conclua sobre os registros reais da MDIA3.'),
('Reprodução','Instale requirements.txt e execute os comandos de reprodução descritos no README.'),
('README','Reúne matrícula, organização do repositório e comandos de reprodução.'),
('IA','Inteligência artificial foi utilizada como apoio à elaboração, com revisão e validação integradas ao trabalho acadêmico.'),
]
table(ws,['Campo','Descrição'],intro)
ws.column_dimensions['C'].width=18
ws.column_dimensions['D'].width=60
ws.print_area='C3:D17'
for row,height in {9:24,10:28,11:24,12:30,13:34,14:34,15:30,16:30,17:42}.items(): ws.row_dimensions[row].height=height

s=wb.create_sheet('Materialidade'); setup(s,'Materialidade','Premissas acadêmicas ilustrativas e análise de sensibilidade','BRL mil; as porcentagens são julgamentos didáticos, não percentuais mandatórios da NBC TA 320')
mat_rows=[
('Resultado antes de IR/CSLL',708270,'DF consolidada 2025, DFC/resultado; 31/12/2025; balanço/resultado, exercício 2025.'),
('Receita líquida',10437548,'DF consolidada 2025, Nota 27; exercício 2025.'),
('Ativo total',12559343,'DF consolidada 2025, balanço patrimonial; 31/12/2025.'),
('Benchmark escolhido',0.05,'Premissa acadêmica: 5% do resultado antes de IR/CSLL, medida potencialmente relevante para entidade lucrativa; PBT caiu 14,0% e inclui subvenções expressivas, portanto foi comparado a receita, ativo e PBT sem a linha de subvenção.'),
('Materialidade execução (% global)',0.50,'Premissa conservadora para cenário de auditoria inicial e riscos elevados; percentual é julgamento acadêmico, não regra normativa.'),
('Limite trivial (% global)',0.05,'Premissa didática: 5% da global.'),
]
for ix,(label,val,src) in enumerate(mat_rows,9):
 s.cell(ix,3,label); s.cell(ix,3).font=Font(color=BLUE); s.cell(ix,3).comment=Comment(src,'Fonte')
 s.cell(ix,4,val); s.cell(ix,4).font=Font(color=BLUE); s.cell(ix,4).comment=Comment(src,'Fonte')
 s.cell(ix,4).number_format='0.0%' if ix in (12,13,14) else '#,##0.000;(#,##0.000);-'
for row,label,formula in [(15,'Materialidade global','=D9*D12'),(16,'Materialidade para execução','=D15*D13'),(17,'Limite trivial','=D15*D14'),(18,'Benchmark alternativo: 0,5% receita','=D10*D34'),(19,'Benchmark alternativo: 1% ativo','=D11*D35'),(20,'Global / receita','=D15/D10'),(21,'Global / ativo','=D15/D11')]:
 s.cell(row,3,label); s.cell(row,4,formula); s.cell(row,4).font=Font(color=BLACK); s.cell(row,4).number_format='0.00%' if row in (20,21) else '#,##0.000;(#,##0.000);-'
s['C23']='Interpretação'; s['C23'].fill=PatternFill('solid',fgColor=LIGHT); s['C24']='Materialidade não é uma conta do ativo nem um corte automático por conta. O PBT foi escolhido provisoriamente por ser medida de desempenho potencialmente relevante para entidade lucrativa, mas caiu 14,0% e inclui BRL 463.419 mil de subvenção estadual. Benchmarks alternativos e a exclusão mecânica dessa linha são sensibilidades, não ajuste contábil nem escolha automática do benchmark correto.'; s.merge_cells('C24:H25'); s['C24'].alignment=Alignment(wrap_text=True,vertical='top')
for row,label,value in [(27,'Subvenção estadual na DRE (2025)',463419),(28,'PBT sem essa linha — sensibilidade','=D9-D27'),(29,'Global a 5% do PBT sensível','=D28*D12'),(30,'Execução a 50%','=D29*D13'),(31,'Trivial a 5% da global','=D29*D14'),(32,'Benchmark alternativo: receita','=D10*D34'),(33,'Benchmark alternativo: ativo','=D11*D35')]:
 s.cell(row,3,label); s.cell(row,4,value); s.cell(row,4).number_format='#,##0.000;(#,##0.000);-'; s.cell(row,3).alignment=Alignment(wrap_text=True); s.cell(row,4).alignment=Alignment(wrap_text=True)
 if isinstance(value,str) and value.startswith('='): s.cell(row,4).font=Font(color=BLACK)
 else: s.cell(row,4).font=Font(color=BLUE); s.cell(row,4).comment=Comment('DFs consolidadas 2025; sensibilidade didática, sem ajuste contábil.','Fonte')
s['C34']='Percentual alternativo sobre receita'; s['D34']=0.005; s['D34'].number_format='0.00%'; s['D34'].font=Font(color=BLUE); s['D34'].comment=Comment('Premissa de sensibilidade didática: 0,5% da receita líquida publicada.','Fonte')
s['C35']='Percentual alternativo sobre ativo'; s['D35']=0.01; s['D35'].number_format='0.00%'; s['D35'].font=Font(color=BLUE); s['D35'].comment=Comment('Premissa de sensibilidade didática: 1,0% do ativo total publicado.','Fonte')
for row,label,formula in [(36,'Execução alternativa: 0,5% receita','=D18*D13'),(37,'Trivial alternativo: 0,5% receita','=D18*D14'),(38,'Execução alternativa: 1% ativo','=D19*D13'),(39,'Trivial alternativo: 1% ativo','=D19*D14')]:
 s.cell(row,3,label); s.cell(row,4,formula); s.cell(row,4).font=Font(color=BLACK); s.cell(row,4).number_format='#,##0.000;(#,##0.000);-'
for col in range(3,9): s.column_dimensions[get_column_letter(col)].width=27

s=wb.create_sheet('Saldos'); setup(s,'Saldos','Principais saldos publicados e composição do imobilizado','BRL mil; saldos consolidados')
rows=[
('Ativo total',12559343,12768867,'Balanço consolidado'),('Imobilizado',3570990,3590675,'Balanço consolidado/Nota 11'),('Intangível',2416590,2414480,'Balanço consolidado/Nota 12'),('Caixa e equivalentes',1888270,2152587,'Balanço consolidado/Nota 5'),('Clientes — total comparável',1843604,1670045,'Balanço consolidado/Nota 6; 2024 inclui circulante + não circulante'),('Estoques',1675464,1687637,'Balanço consolidado/Nota 7'),('Passivo circulante',2235103,2732737,'Balanço consolidado'),('Receita líquida',10437548,9662833,'Nota 27'),('Resultado antes IR/CSLL',708270,823185,'Demonstração do resultado/fluxo de caixa'),('Custo bruto imobilizado',7205695,None,'Nota 11'),('Depreciação acumulada',3634705,None,'Nota 11'),
]
table(s,['Rubrica','2025','2024','Referência'],[(a,b,c,d) for a,b,c,d in rows],8,3,'Saldos')
for i in range(9,9+len(rows)):
 for j in [4,5]: s.cell(i,j).number_format='#,##0.000;(#,##0.000);-'
comp=[('Edificações',3000,1740776,548812,1191964),('Máquinas e equipamentos',210000,3682384,2388537,1293847),('Móveis e utensílios',160000,145342,97227,48115),('Veículos',10000,256948,127743,129205),('Instalações',40000,572492,306773,265719),('Imobilizações em andamento',30000,415229,2220,413009),('Outros',47000,392524,163393,229131)]
start=23
for j,h in enumerate(['Classe','Registros sintéticos','Custo bruto','Depreciação acumulada','Saldo líquido'],3): s.cell(start,j,h).fill=PatternFill('solid',fgColor=DARK); s.cell(start,j).font=Font(color=WHITE,bold=True); s.cell(start,j).alignment=Alignment(wrap_text=True)
for i,row in enumerate(comp,start+1):
 for j,v in enumerate(row,3):
  c=s.cell(i,j,v); c.font=Font(color=BLUE); c.comment=Comment('Fonte saldo: DFs anuais completas, Nota 11, páginas impressas 60–62, consolidado, 31/12/2025. Número de linhas: regra didática do script, não fonte da companhia.','Fonte')
  if j>=5: c.number_format='#,##0.000;(#,##0.000);-'
for j,v in enumerate(['Total','=SUM(D24:D30)','=SUM(E24:E30)','=SUM(F24:F30)','=SUM(G24:G30)'],3):
 s.cell(31,j,v); s.cell(31,j).font=Font(bold=True,color=BLACK if isinstance(v,str) and v.startswith('=') else BLUE); s.cell(31,j).border=Border(top=Side(style='thin',color=BLACK))
for j in range(5,8): s.cell(31,j).number_format='#,##0.000;(#,##0.000);-'
s.column_dimensions['G'].width=22

s=wb.create_sheet('Seleção da conta'); setup(s,'Seleção da conta para a base','Comparação de magnitude e fatores de risco dos principais saldos de ativo','Saldos consolidados 2025; magnitude não é sinônimo de materialidade normativa')
selection_rows=[
('Imobilizado','=Saldos!D10','=D9/Saldos!D9','Maior linha individual de ativo; existência, direitos, baixas, depreciação, projetos em andamento e impairment.','Selecionada para a base obrigatória pela maior magnitude individual entre as principais rubricas e pelos riscos inerentes.'),
('Intangível','=Saldos!D11','=D10/Saldos!D9','Estimativa/impairment de goodwill e marcas; alto julgamento, apesar de saldo inferior ao imobilizado.','Risco de estimativa elevado; não é a maior linha individual por saldo.'),
('Caixa e equivalentes','=Saldos!D12','=D11/Saldos!D9','Existência, restrição, direitos e confirmação bancária.','Saldo relevante; testes dependem de confirmações e extratos.'),
('Clientes','=Saldos!D13','=D12/Saldos!D9','Existência, direitos, aging e perda esperada.','Saldo relevante; expectativa analítica requer investigação.'),
('Estoques','=Saldos!D14','=D13/Saldos!D9','Existência física, custo, obsolescência e valor realizável líquido.','Saldo relevante; exige observação física e testes de custo.')]
table(s,['Conta','Saldo 2025 (BRL mil)','% do ativo','Fatores de risco','Conclusão acadêmica'],selection_rows)
for row in range(9,9+len(selection_rows)): s.cell(row,5).number_format='0.00%'
for row in range(9,14): s.row_dimensions[row].height=75
s['C17']='Conclusão'; s['C17'].fill=PatternFill('solid',fgColor=LIGHT); s['C18']='A NBC TA 320 não determina uma rubrica “mais material” apenas pelo saldo. Para cumprir a exigência de base, o imobilizado é a escolha mais defensável entre as principais linhas individuais: é a maior por valor (BRL 3.570.990 mil; 28,43% do ativo) e possui riscos inerentes relevantes. A avaliação de risco também destaca intangível e receita, e a hierarquia pode mudar com informação interna ou necessidades específicas de usuários.'; s.merge_cells('C18:H20'); s['C18'].alignment=Alignment(wrap_text=True,vertical='top')
s.row_dimensions[18].height=35; s.row_dimensions[19].height=25; s.row_dimensions[20].height=25

s=wb.create_sheet('Riscos'); setup(s,'Matriz de riscos','Riscos inerentes, afirmações e resposta planejada','Planejamento; não equivale a procedimento executado em dados reais')
risk=[
('Receita líquida 10.437.548','Alto','Ocorrência, exatidão, corte, apresentação','Fraude presumida no reconhecimento; venda faturada sem transferência de controle','Testar entregas de dezembro/janeiro, devoluções, contratos, lançamentos e razão','Base de vendas/entrega e razão indisponíveis'),
('Imobilizado 3.570.990','Alto','Existência, direitos, integridade, valoração, corte, classificação','Adições/baixas, vida útil/depreciação, projetos em andamento, impairment','Auxiliar ao razão, títulos, inspeção, recálculo, baixas, capex em andamento','Só saldos públicos e população sintética'),
('Intangível 2.416.590','Alto','Existência, direitos, valoração, classificação, divulgação','Goodwill/marcas sem amortização e fluxo descontado','Testar CGUs, budgets, forecasts, taxa 13,04%, crescimento terminal 3,9%, sensibilidade','Modelo detalhado ausente'),
('Clientes 1.843.604; ECL 60.051','Alto/moderado','Existência, direitos, valoração','Aging, exposição vencida, premissas ECL','Confirmações, recebimentos subsequentes, recalcular aging e modelo','Auxiliar e confirmações ausentes'),
('Estoques 1.675.464; perda 31.481','Alto','Existência, integridade, valoração, corte','Contagem, custo médio, obsolescência/VRL','Contagem observada, teste de custos, venda subsequente','Sem contagem e fichas de custo'),
('Provisões 200.718; contingências possíveis 2.307.733','Alto qualitativo','Integridade, mensuração, apresentação','Completude e classificação probabilidade','Carta jurídica, autos, atualização e cálculo','Sem carta/autos; contingência possível não é passivo reconhecido'),
('Caixa 1.888.270; derivativos','Moderado/alto','Existência, direitos, valor justo, classificação','Restrição, contraparte, hedge','Confirmações, extratos, recálculo e documentação hedge','Sem confirmações'),
('Dívida 1.419.965; arrendamentos 322.005','Moderado','Integridade, obrigação, valoração, classificação','Vencimentos, covenant, taxa, juros','Confirmações, contratos, recálculo, busca de passivo','Sem contratos completos'),
]
table(s,['Área','Risco','Afirmações','Fator','Resposta planejada','Limitação'],risk)

s=wb.create_sheet('Amostra'); setup(s,'Amostragem','Desenho didático sobre a população sintética','Sem extrapolação para MDIA3; AAS de atributo não é amostra monetária')
rows=[('População',500000,'Itens sintéticos com IDs únicos; não são cadastro da companhia'),('Unidade','1 linha = 1 item fictício','Não equivale a unidade real de cadastro'),('AAS descritiva',384,'Sem reposição; seed 5302025; objetivo: presença de ID, classe e 3 campos monetários; 0/384 ausentes'),('Seleção dirigida','25 maiores saldos + censo das flags','Não probabilística; não projetável'),('Amostra combinada',441,'AAS mais seleções dirigidas, IDs deduplicados'),('MUS sistemática','606 unidades monetárias; 603 IDs únicos','Distorção tolerável proxy = PM BRL 17.706,75 mil; fator 3,0; erro esperado zero'),('Intervalo MUS','=Saldos!D10/606','Início aleatório reproduzível; consultar amostra_monetaria_MUS.csv'),('Projeção pontual MUS','BRL 41.785,446 mil','BRL 30.000 mil em três itens de certeza + dois itens × intervalo; estimativa pontual didática, não limite superior nem projeção para MDIA3.'),('Detalhe da projeção','pt530_detalhe_projecao_mus.csv','Cinco baixas sintéticas selecionadas; pontos MUS, saldo, certeza, tainting e parcela projetada.'),('Flags do cenário','37 exceções','5 baixas; 12 projetos; 12 etiquetas; 8 não localizados'),('Limitação','Desenho didático em base sintética','Projeção e atributos não sustentam opinião; recalcular com população real em auditoria')]
table(s,['Parâmetro','Valor','Comentário'],rows)
s['D15'].number_format='#,##0.000;(#,##0.000);-'

s=wb.create_sheet('Exceções'); setup(s,'Distorções e exceções','Separar cenário simulado de achados reais','Nenhum item desta planilha é ocorrência real da companhia')
exrows=[('Baixas não registradas',5,39000,'1,10× global; 2,20× execução; 5,51% PBT','Factual apenas no cenário; custo 48.500, depreciação 9.500, líquido 39.000. Correção suposta sem produto da alienação reduziria o resultado antes de tributos em 39.000.'),('Itens em andamento possivelmente disponíveis',12,None,'Não quantificado','Inconsistência sintética de dados; efeito de depreciação depende de data real, custo, vida útil e valor residual.'),('Etiqueta ausente',12,None,'Não quantificado','Controle/rastreabilidade; não é automaticamente distorção contábil.'),('Não localizado simulado',8,3.848,'Equivale a BRL 3.848 correntes','Resultado fabricado; não é inspeção real; sem evidência para ajuste.'),('Distorções reais identificadas',0,0,'Não testadas','Nenhuma distorção real pode ser declarada sem evidência dos registros da companhia.')]
table(s,['Classe','Qtd','Exposição potencial (BRL mil)','Comparação','Conclusão/limite'],exrows)

s=wb.create_sheet('Analíticos'); setup(s,'Procedimentos analíticos','Expectativas, limiar de investigação e resultados reproduzíveis','BRL mil; expectativas didáticas baseadas em dados públicos agregados')
analytic_rows=[
('Clientes: expectativa proporcional à receita','=Saldos!D13','=Saldos!E13*(Saldos!D16/Saldos!E16)','=D9-E9','=Materialidade!D16','2024 total comparável BRL 1.670.045; expectativa 2025 BRL 1.803.940,402; diferença BRL 39.663,598 (~2,24x PM). Investigar mix, prazo, aging, moeda, concentração, recebimentos e ECL; não é erro identificado.'),
('Receita líquida: variação anual','=Saldos!D16','=Saldos!E16','=Saldos!D16/Saldos!E16-1','=Materialidade!D16','Aumento percentual: decompor por mês/produto/preço/volume e testar corte; diferença percentual não é erro monetário.'),
('Resultado antes de IR/CSLL: variação anual','=Saldos!D17','=Saldos!E17','=Saldos!D17/Saldos!E17-1','=Materialidade!D16','Queda de resultado superior à mudança de receita; testar margem, financeiro, subvenções e itens não recorrentes.'),
('Imobilizado líquido: variação anual','=Saldos!D10','=Saldos!E10','=Saldos!D10/Saldos!E10-1','=Materialidade!D16','Conciliar adições, depreciação, baixas, transferências, impairment e direitos de uso.'),
]
table(s,['Analítico','Saldo 2025','Base/expectativa 2024','Diferença/variação','Limiar de triagem','Investigação/resposta'],analytic_rows)
for row in range(9,9+len(analytic_rows)):
 for col in (4,5,7): s.cell(row,col).number_format='#,##0.000;(#,##0.000);-'
 s.cell(row,6).number_format='#,##0.000;(#,##0.000);-' if row==9 else '0.0%;(0.0%);-'
for row in range(9,9+len(analytic_rows)): s.row_dimensions[row].height=58
s=wb.create_sheet('Programa'); setup(s,'Programas de auditoria','Procedimentos por afirmação, natureza, época e extensão','Planejados para cenário de auditoria; execução real não ocorreu')
program_rows=[
('PT240/330','Receita: ocorrência, exatidão, corte; fraude presumida','Teste substantivo de notas, pedidos, prova de entrega/aceite, devoluções subsequentes e lançamentos manuais; conciliar faturamento ao razão.','Detalhe + analítico','Dezembro 2025 e janeiro/fevereiro 2026; incluir seleção antes/depois do corte.','Estratificar por mês/produto e risco; tamanho após tolerância e risco definidos.','Não executado: não há base de vendas nem documentos.'),
('PT315/330','SAP S/4HANA e dados de reporte: integridade, exatidão, acesso','Entender migração Simplifique concluída em 2024, integração declarada de 47 sistemas; testar controles gerais de TI, interfaces, acessos e lançamentos manuais.','Walkthrough + testes de controles + substantivo','Migração 2024 e operação 2025; controles ao longo do exercício.','Selecionar interfaces críticas, usuários privilegiados, mudanças e journals por risco.','Risco identificado em divulgação; sem acesso a SAP, logs ou ITGC.'),
('PT530','Imobilizado: existência, valor e baixas','Seleção MUS sistemática da população sintética; inspeção bidirecional apenas se ativos/evidências reais forem disponibilizados.','Amostragem monetária + detalhe','Data-base 31/12/2025; altas e baixas no exercício.','606 unidades monetárias; início aleatório e intervalo no script/CSV; PM usada como proxy de distorção tolerável.','Executado somente no conjunto sintético; sem inferência para a MDIA3.'),
('PT330','Imobilizado: direitos, mensuração e depreciação','Vouching de adições a nota/contrato; recálculo de depreciação; exame de baixas, alienações e transferência de projetos.','Detalhe','Amostra de adições/baixas de 2025 e itens de maior valor.','Itens selecionados pelo MUS e itens de alto risco; cobertura orientada por afirmação.','Não executado em registros da empresa; classe de direito de uso não separada na base sintética.'),
('PT501','Estoques: existência e valor realizável líquido','Acompanhar inventário físico; testar contagens, custo médio, obsolescência e venda subsequente.','Observação + detalhe','Data do inventário e testes de corte no encerramento.','Locais/materialidade por local; itens de alto valor e itens selecionados.','Sem contagem, fichas de custo ou vendas posteriores granulares.'),
('PT505/330','Disponibilidades: existência, direitos, restrições e classificação','Confirmar bancos sob controle do auditor; reconciliar extratos, aplicações, restrições e corte.','Confirmação externa + detalhe','Próximo a 31/12/2025 e conciliações posteriores.','População bancária real e todas as contas relevantes; valores/contatos ainda não disponíveis.','Somente modelo de confirmação; sem população, envio, resposta ou teste.'),
('PT505/330','Passivos financeiros: completude, obrigação, mensuração e covenant','Confirmar financiadores; inspecionar contratos, juros, vencimentos, garantias e covenants; recalcular saldos.','Confirmação externa + detalhe','Data-base e pagamentos posteriores.','Todos os financiadores relevantes e contratos; seleção após obter razão/contratos.','Sem confirmação, contratos completos ou razão.'),
('PT330/550','Derivativos: existência, direitos, valor justo e divulgação','Confirmar termos com contrapartes; recalcular valor justo e avaliar designação/efetividade de hedge.','Confirmação externa + recálculo','Posições em 31/12/2025 e transações subsequentes.','Censo das posições relevantes; suporte/contratos ausentes.','Somente divulgações agregadas; sem confirmação ou recálculo independente.'),
('PT250/501','Provisões/contingências: completude e mensuração','Solicitar carta jurídica, processos, cálculo de provisões e avaliação de probabilidade; verificar CPC 25 e divulgação.','Confirmação jurídica + inspeção','Até a data do relatório.','Censo de processos materiais; população jurídica não disponível.','Sem carta, autos, assessoria jurídica ou razão auxiliar.'),
('PT505','Clientes: existência/direitos e perda esperada','Confirmação positiva de saldos, recebimentos subsequentes, aging e recálculo da ECL.','Confirmação externa + detalhe','Confirmações próximas a 31/12/2025; alternativas após não respostas.','Estratificar por saldo, atraso, disputa e risco; amostra definida após população real.','Minuta apenas; nenhuma confirmação foi enviada.'),
('PT540','Goodwill e marcas: valoração','Testar CGUs, modelo DCF, dados, projeções, taxa de desconto, crescimento terminal, sensitividade e viés.','Teste de estimativa + analítico','Teste anual e orçamento aprovado para 2026–2030.','Cobertura de todas as CGUs relevantes; cenário de taxa/growth não substitui reexecução do modelo.','Somente sensibilidade relativa calculada; modelo/folga ausentes.'),
('PT550/560','Partes relacionadas e eventos subsequentes: completude/apresentação','Cruzar cadastro, atas, contratos, pagamentos, publicações e eventos até data do relatório; avaliar CPC 05/CPC 24.','Indagação + inspeção + analítico','Do encerramento a 26/02/2026.','Censo de atas/eventos e seleção de transações/saldos materiais.','Fontes públicas; não houve indagação, atas privadas ou razão.'),
('PT570','Continuidade: apresentação/divulgação','Testar fluxo de caixa projetado, covenants, vencimentos, linhas de crédito, capex e cenários de estresse.','Analítico + detalhe','Pelo menos 12 meses da data de autorização.','Todos os financiamentos e covenants relevantes; cenários adversos de liquidez.','Não executado; previsões e contratos completos ausentes.'),
('PT600','Consolidação e componentes: integridade/valoração','Validar perímetro, balancetes, eliminações, políticas, câmbio e trabalho de auditores de componentes; CPC 15/36.','Consolidação + revisão de componentes','Fechamento anual e saldos comparativos.','Cobrir componentes significativos e riscos por componente; extensão após materialidade de grupo/componente.','Mapeamento só documental; sem balancetes e papéis dos componentes.'),
]
table(s,['PT/área','Risco/afirmação','Procedimento','Natureza','Época','Extensão/amostra','Execução/limite'],program_rows)
for row in range(9,9+len(program_rows)): s.row_dimensions[row].height=105
s.row_breaks.append(Break(id=15))

s=wb.create_sheet('Normas'); setup(s,'Normas','Versões, vigência e status da evidência','NBC TA 200–600: escopo principal; normas 700+ apenas como referência suplementar ao modelo')
with open(MATRIX,encoding='utf-8-sig',newline='') as f:
    matrix=list(csv.DictReader(f,delimiter=';'))
norms_with_vigencia=[]
for r in matrix:
    obj=r['Procedimento e evidência']
    if len(obj)>240: obj=obj[:237]+'...'
    norms_with_vigencia.append((r['PT'],r['Norma/versão indicada no compêndio de referência'],r['Vigência conforme cláusula do compêndio de referência; aplicabilidade a 2025'],obj,r['Status/limitação']))
table(s,['PT','Norma/versão','Vigência/base','Objeto/procedimento','Status/limite'],norms_with_vigencia)
for row in range(9,9+len(norms_with_vigencia)): s.row_dimensions[row].height=55

# Source comments, consistent typography, and finance-standard color semantics
for sh in wb.worksheets:
    sh.sheet_view.zoomScale=85
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None:
                if c.column>=3 and c.row>=8: c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.data_type=='f' or (isinstance(c.value,(int,float)) and not isinstance(c.value,bool)):
                    fmt=c.number_format
                    if fmt=='General': fmt='#,##0.###;(#,##0.###);-'
                    if not fmt.startswith('[$-416]'): c.number_format='[$-416]'+fmt
                font=copy(c.font); font.name='Arial'
                if c.data_type=='f':
                    font.color=GREEN if '!' in str(c.value) else BLACK
                elif isinstance(c.value,(int,float)) and not isinstance(c.value,bool):
                    if c.comment is None: c.comment=Comment(source,'Fonte/assunção')
                    font.color=BLUE
                c.font=font
    sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.print_title_rows='8:8'
wb.save(OUT)
wb.properties.creator='Trabalho acadêmico — matrícula 202062044'
wb.properties.title='Papéis de trabalho de auditoria didática — MDIA3 2025'
wb.properties.subject='Materialidade, riscos, amostragem sintética e conciliação; sem opinião independente'
wb.save(OUT)
# open and confirm file structure and critical formulas
check=load_workbook(OUT,data_only=False)
assert check['Materialidade']['D15'].value=='=D9*D12'
assert len(check.sheetnames)>=7
print(f'Criado {OUT}; abas={check.sheetnames}; fórmula materialidade={check["Materialidade"]["D15"].value}')
