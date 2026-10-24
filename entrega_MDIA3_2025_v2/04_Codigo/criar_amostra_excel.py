#!/usr/bin/env python3
from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

SCRIPT_DIR=Path(__file__).resolve().parent
PROJECT=SCRIPT_DIR.parent if SCRIPT_DIR.name=='04_Codigo' else SCRIPT_DIR
DATA_DIR=PROJECT/'02_Base_e_amostras' if (PROJECT/'02_Base_e_amostras').exists() else PROJECT
SHEET_DIR=PROJECT/'03_Planilhas' if (PROJECT/'03_Planilhas').exists() else PROJECT
SOURCE=DATA_DIR/'amostra_combinada_auditoria.csv'
OUT=SHEET_DIR/'amostra_imobilizado_sintetica_legivel.xlsx'
wb=Workbook(); ws=wb.active; ws.title='Amostra 441'
ws.sheet_view.showGridLines=False
with SOURCE.open('r',encoding='utf-8-sig',newline='') as f:
    reader=csv.DictReader(f,delimiter=';')
    headers=reader.fieldnames
    for col,h in enumerate(headers,1):
        c=ws.cell(1,col,h); c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='135B44'); c.alignment=Alignment(wrap_text=True,vertical='center')
    for rownum,row in enumerate(reader,2):
        for col,h in enumerate(headers,1):
            val=row.get(h,'')
            if h.endswith('_R$_mil') and val:
                # Preserve as numeric Excel cell. Data carries three decimal places of R$ mil.
                val=float(val.replace('.','').replace(',','.'))
            c=ws.cell(rownum,col,val); c.alignment=Alignment(vertical='top',wrap_text=True)
            if h.endswith('_R$_mil') and val!='': c.number_format='#,##0.000;(#,##0.000);-'
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
for col in range(1,ws.max_column+1):
    letter=get_column_letter(col)
    maxlen=min(42,max(12,max(len(str(ws.cell(r,col).value or '')) for r in range(1,min(ws.max_row,70)+1))+2))
    ws.column_dimensions[letter].width=maxlen
ws.row_dimensions[1].height=45
ws.insert_rows(1,2)
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ws.max_column)
ws.cell(1,1,'AMOSTRA SINTÉTICA DE IMOBILIZADO — MATRÍCULA 202062044').font=Font(size=14,bold=True,color='FFFFFF')
ws.cell(1,1).fill=PatternFill('solid',fgColor='135B44')
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ws.max_column)
ws.cell(2,1,'Somente dados fictícios; os saldos agregados foram calibrados à Nota 11, mas as linhas não são ativos reais da MDIA3. A seleção combina AAS, direcionamento e censo de flags.').alignment=Alignment(wrap_text=True)
ws.cell(2,1).comment=Comment('Aviso de escopo: itens e atributos gerados artificialmente; não constituem cadastro patrimonial da companhia.','Fonte')
ws.freeze_panes='A4'; ws.auto_filter.ref=f'A3:{get_column_letter(ws.max_column)}{ws.max_row}'
ws.sheet_view.zoomScale=75
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.oddFooter.center.text='Amostra sintética | Página &P de &N'
wb.properties.title='Amostra sintética de imobilizado - exercício acadêmico'
wb.properties.subject='Dados simulados; matrícula 202062044'
wb.save(OUT)
check=load_workbook(OUT,read_only=True,data_only=True)
assert check['Amostra 441'].max_row==444, check['Amostra 441'].max_row
print(f'Criado {OUT}; linhas de dados=441; linhas totais com cabeçalho e avisos={check["Amostra 441"].max_row}')
